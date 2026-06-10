from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\test-flow-app")
OUT = ROOT / "reports" / "planning-scenarios-report.docx"
SCREENSHOTS = ROOT / "reports" / "screenshots"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa=9360, indent_dxa=120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    run.font.size = Pt(9)
    run.font.italic = True
    run.font.color.rgb = RGBColor(85, 85, 85)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_width(table)
    for idx, text in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = text
        set_cell_shading(cell, "F2F4F7")
        set_cell_width(cell, widths[idx])
        set_cell_margins(cell)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    for row in rows:
        cells = table.add_row().cells
        for idx, text in enumerate(row):
            cells[idx].text = text
            cells[idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cells[idx], widths[idx])
            set_cell_margins(cells[idx])
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(8.5)
    doc.add_paragraph()


def add_code(doc: Document, title: str, file_ref: str, code: str) -> None:
    h = doc.add_paragraph(style="Heading 3")
    h.add_run(title)
    p_ref = doc.add_paragraph()
    p_ref.paragraph_format.space_after = Pt(3)
    r = p_ref.add_run(file_ref)
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(85, 85, 85)
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_width(table, width_dxa=9360, indent_dxa=120)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, "F7F9FB")
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(code.strip())
    run.font.name = "Courier New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    run.font.size = Pt(7.5)
    doc.add_paragraph()


def add_picture(doc: Document, filename: str, caption: str) -> None:
    path = SCREENSHOTS / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(6.25))
    add_caption(doc, caption)


def setup_document() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    return doc


def build() -> None:
    doc = setup_document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(3)
    run = title.add_run("Отчет о сценариях планирования строительной техники")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    srun = subtitle.add_run(
        "Приложение для мониторинга и планирования использования строительной техники на предприятии дорожной отрасли"
    )
    srun.font.size = Pt(12)
    srun.font.color.rgb = RGBColor(85, 85, 85)

    meta_rows = [
        ["Дата подготовки", date.today().strftime("%d.%m.%Y")],
        ["Проверенная ветка", "feature/diploma-theme-refactor"],
        ["Ключевой модуль", "План-график: объекты, этапы, потребность, назначения, ТО/ремонтные риски"],
        ["Автотесты", "equipment-plans.service.spec.ts: 28/28; service-events.service.spec.ts: 4/4"],
    ]
    add_table(doc, ["Параметр", "Значение"], meta_rows, [2200, 7160])

    doc.add_heading("1. Назначение отчета", level=1)
    doc.add_paragraph(
        "Цель отчета - показать, что функция планирования не является статической таблицей. "
        "Backend рассчитывает технологические этапы дорожного объекта, потребность в технике, "
        "доступность машин по датам и сменам, а также реакцию на ТО и ремонты. Frontend отображает "
        "результат в виде выбора объекта, диаграммы Ганта, назначений техники и карточек рисков."
    )
    doc.add_paragraph(
        "В текущей версии область планирования не использует роли водителей, механиков или модераторов. "
        "Актуальная модель доступа для отчета: администратор и пользователь. Пользователь работает с "
        "планированием и мониторингом, администратор дополнительно управляет пользователями."
    )

    doc.add_heading("2. Архитектурная схема планирования", level=1)
    diagram = (
        "RoadWorkTypeTemplate -> RoadWorkStageTemplate -> RoadWorkStageEquipmentTemplate\n"
        "                         |                         |\n"
        "                         v                         v\n"
        "ConstructionSite -> RoadWorkStage -> EquipmentDemand -> EquipmentPlanAssignment -> FleetVehicle\n"
        "                                                             ^\n"
        "FleetRepairTemplate -> FleetServiceEvent --------------------|\n"
        "   (тип техники, длительность)       (ТО/ремонтное окно блокирует назначение)"
    )
    p = doc.add_paragraph()
    r = p.add_run(diagram)
    r.font.name = "Courier New"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    r.font.size = Pt(9)

    doc.add_paragraph(
        "Виды работ и этапы выступают справочниками. При расчете объект получает собственный набор этапов, "
        "потребностей и сменных назначений. ТО и ремонты не являются отдельной справкой без влияния на план: "
        "их даты блокируют технику и могут вызвать переназначение или сдвиг этапов."
    )

    doc.add_heading("3. Матрица сценариев планирования", level=1)
    scenario_rows = [
        ["1", "Объект не выбран", "Пользователь открыл План-график", "Графики не отображаются, предлагается выбрать объект", "UI screenshot 1"],
        ["2", "Объект без плана", "Есть дорожный объект без этапов", "Показывается кнопка 'Спланировать'", "UI screenshot 2"],
        ["3", "Первичный расчет", "Выбран объект, вид работ, дата начала, длина, ширина, плечо доставки", "Создается черновик этапов и потребностей", "generateDraft DTO + wizard"],
        ["4", "Автоматическое назначение", "Есть свободная техника нужного типа", "Создаются сменные назначения на даты этапа", "applyDraft, строки назначения"],
        ["5", "Занятость на другом объекте", "Техника уже стоит в плане на ту же дату/смену", "Она исключается из доступных; ручной выбор отклоняется", "getOccupiedVehicleIds + tests"],
        ["6", "Занятость внутри текущего черновика", "Две перекрывающиеся стадии требуют одну и ту же машину", "Вторая стадия видит конфликт и дефицит", "draftOccupiedSlotKeys + test"],
        ["7", "Техника в ТО/ремонте до сохранения", "Есть сервисное окно на дату этапа", "Машина считается недоступной, этап сдвигается или возникает дефицит", "serviceBlockingWhere + test"],
        ["8", "Нехватка техники", "Требуемое количество больше доступного", "Формируется высокий риск и дефицит покрытия", "buildDraftRisk + coverage"],
        ["9", "Автосдвиг этапов", "autoSchedule=true и текущие даты заняты", "Backend ищет ближайший день, где все правила этапа покрываются", "autoSchedule loop + test"],
        ["10", "Сохранение плана", "Черновик подтвержден пользователем", "Транзакционно создаются этапы, потребности и назначения", "applyDraft transaction"],
        ["11", "Перерасчет объекта", "У объекта уже есть план", "Старые этапы/потребности/назначения объекта удаляются и строятся заново; свои назначения не блокируют перерасчет", "replaceExisting + test"],
        ["12", "Ремонт после сохранения", "По назначенной машине создан ремонт на даты плана", "Появляется сервисный риск в карточке объекта", "getServiceRiskResolution + UI"],
        ["13", "Переназначение", "Есть свободная машина того же типа", "Пользователь выбирает замену; backend переносит смены на новую технику", "resolveServiceRisk replace"],
        ["14", "Нет свободной замены", "Все машины того же типа заняты или в ремонте", "Система предлагает сдвиг этапов после ремонта", "resolveServiceRisk shift"],
        ["15", "Конфликт при сохранении", "Параллельно появилась занятость той же машины", "Уникальный индекс БД превращается в понятную ошибку планирования", "Prisma P2002 + test"],
        ["16", "Невалидный этап", "Передан тип этапа вне enum", "DTO отклоняет запрос с ошибкой 400", "class-validator"],
        ["17", "Сброс плана", "У объекта есть сохраненный план", "Удаляются назначения, потребности и этапы объекта", "resetSitePlan"],
        ["18", "Жизненный цикл смены", "Назначение сохранено", "Статус смены может перейти planned -> in_progress -> completed/failed, фиксируются плановые и фактические часы", "EquipmentPlanAssignment model"],
        ["19", "План без факта", "Дата работ прошла, actualHours не заполнены", "Система может подсветить отсутствие фактических часов как мониторинговый риск", "EquipmentPlanStats"],
        ["20", "Закрытая заявка сервиса", "Ремонт завершен", "Риск больше не отображается, статус техники восстанавливается", "frontend risk filter + ServiceEventsService tests"],
    ]
    add_table(doc, ["№", "Сценарий", "Предусловие", "Результат системы", "Доказательство"], scenario_rows, [450, 1750, 2350, 3000, 1810])

    doc.add_heading("4. Демонстрационные экраны", level=1)
    add_picture(
        doc,
        "01-planning-no-object-crop.png",
        "Скриншот 1. До выбора объекта диаграммы и назначения не отображаются.",
    )
    add_picture(
        doc,
        "02-empty-site-plan-crop.png",
        "Скриншот 2. Объект без плана показывает единственное действие - спланировать.",
    )
    add_picture(
        doc,
        "06-wizard-input.png",
        "Скриншот 3. Мастер планирования: ввод исходных данных - вид работ, дата начала и протяженность.",
    )
    add_picture(
        doc,
        "07-wizard-stages.png",
        "Скриншот 4. Мастер планирования: этапы берутся из шаблона вида работ, но могут быть изменены для конкретного расчета.",
    )
    add_picture(
        doc,
        "08-wizard-equipment.png",
        "Скриншот 5. Мастер планирования: подбор свободной техники, отображение дефицита и кнопка автоматического сдвига этапов.",
    )
    add_picture(
        doc,
        "09-wizard-final.png",
        "Скриншот 6. Мастер планирования: итоговая диаграмма и состав назначенной техники перед сохранением.",
    )
    add_picture(
        doc,
        "03-gantt-saved-plan-crop.png",
        "Скриншот 7. Сохраненный план: этапы дорожных работ выведены на диаграмму Ганта, ниже - сменные назначения.",
    )
    add_picture(
        doc,
        "04-service-risk-replacement-crop.png",
        "Скриншот 8. Ремонт назначенной машины пересекся с планом; система предлагает свободную технику того же типа.",
    )
    add_picture(
        doc,
        "05-service-risk-shift-crop.png",
        "Скриншот 9. Свободной техники того же типа нет; система предлагает сдвинуть этапы после окончания ремонта.",
    )

    doc.add_heading("4.1. Ручная проверка демонстрационных сценариев через API", level=2)
    manual_rows = [
        [
            "GET /equipment-plans/service-risks/:id",
            "Риск с заменой",
            "Geely Atlas Pro 2164 TA-3",
            "2 смены",
            "2 варианта",
            "replace",
            "В карточке риска отображается выбор свободной техники того же типа.",
        ],
        [
            "GET /equipment-plans/service-risks/:id",
            "Риск без замены",
            "МАЗ 5337A2-340P 3098 AM-3",
            "2 смены",
            "0 вариантов",
            "shift, +3 дня",
            "В карточке риска отображается предложение сдвинуть этапы до окончания ремонта.",
        ],
    ]
    add_table(
        doc,
        ["API", "Сценарий", "Назначенная техника", "Затронуто", "Свободные замены", "Рекомендация", "Итог"],
        manual_rows,
        [1450, 1180, 1650, 920, 1080, 1020, 2060],
    )

    doc.add_heading("5. Алгоритм планирования", level=1)
    algorithm_steps = [
        "Пользователь выбирает дорожный объект и запускает расчет.",
        "Backend получает вид работ, дату начала, протяженность, ширину, сменность и плечо доставки.",
        "По виду работ выбираются шаблонные этапы и правила техники; пользователь может изменить этапы только для конкретного расчета.",
        "Для каждого этапа рассчитывается требуемое количество техники и плановые часы.",
        "Backend строит множество занятых слотов: существующие планы других объектов, сервисные окна ТО/ремонтов и уже зарезервированные машины в текущем черновике.",
        "Если включен автосдвиг, система перебирает ближайшие даты и выбирает первый период, где этап можно обеспечить техникой.",
        "После подтверждения расчет сохраняется транзакционно: этапы, потребности и сменные назначения.",
        "Если позже появляется ремонт назначенной техники, frontend запрашивает backend-варианты: заменить техникой того же типа или сдвинуть затронутые этапы.",
    ]
    for step in algorithm_steps:
        doc.add_paragraph(step, style=None)

    doc.add_heading("6. Кодовые доказательства", level=1)
    doc.add_paragraph(
        "Ниже приведены сокращенные фрагменты кода: оставлены ключевые строки и условия, "
        "а не полные тела функций. Абсолютные пути и стартовые строки указаны для быстрой проверки в проекте."
    )
    add_code(
        doc,
        "6.1. Слот планирования и блокирующие сервисные окна",
        r"D:\test-flow-app\backend\src\equipment-plans\equipment-plans.service.ts:409",
        """
409: private planSlotKey(vehicleId: string, workDate: Date, shift: EquipmentPlanShift) {
414:   return `${vehicleId}:${this.dateKey(workDate)}:${shift}`;
417: private serviceBlockingTypes = ['maintenance', 'inspection', 'repair', 'diagnostics'] as const;
424: private serviceBlockSlotKeys(vehicleId: string, startDate: Date, endDate: Date) {
429:   return this.daysBetweenInclusive(startDate, endDate).flatMap((date) => [
430:     this.planSlotKey(vehicleId, date, 'day'),
431:     this.planSlotKey(vehicleId, date, 'night'),
432:   ]);
        """,
    )
    add_code(
        doc,
        "6.2. Исключение занятой техники из доступных вариантов",
        r"D:\test-flow-app\backend\src\equipment-plans\equipment-plans.service.ts:1776",
        """
1776: const getOccupiedVehicleIds = (vehicleType, stageDateKeys, draftOccupiedSlotKeys) => {
1781:   const occupiedByPlans = existingPlans
1788:   const occupiedByService = blockingServices
1800:   const occupiedByDraft = vehicles
1810:   return [...new Set([...occupiedByPlans, ...occupiedByService, ...occupiedByDraft])];
1839: const getAvailableVehicles = (vehicleType, occupiedVehicleIds) => {
1844:   return vehicles.filter((vehicle) =>
1847:     vehicle.type === vehicleType &&
1848:     (vehicle.status === 'active' || vehicle.status === 'reserve') &&
1849:     !occupiedVehicleIdSet.has(vehicle.id));
        """,
    )
    add_code(
        doc,
        "6.3. Автоматический сдвиг этапа до ближайшего доступного окна",
        r"D:\test-flow-app\backend\src\equipment-plans\equipment-plans.service.ts:1909",
        """
1909: if (dto.autoSchedule) {
1912:   for (let delayDays = accumulatedAutoDelayDays; delayDays <= 90; delayDays += 1) {
1917:     const candidateStart = this.addDays(startDate, originalOffsetDays + delayDays);
1922:     if (canCoverStage(stage, candidateStart, draftOccupiedSlotKeys)) {
1923:       accumulatedAutoDelayDays = delayDays;
1924:       effectiveOffsetDays = originalOffsetDays + delayDays;
1925:       break;
1926:     }
1927:   }
        """,
    )
    add_code(
        doc,
        "6.4. Сохранение назначений и запрет ручного выбора занятой техники",
        r"D:\test-flow-app\backend\src\equipment-plans\equipment-plans.service.ts:2233",
        """
2233: const busySelected = sourceVehicles.filter((vehicle) =>
2234:   occupied.has(this.planSlotKey(vehicle.id, workDate, 'day')));
2236: if (busySelected.length > 0) {
2243:   throw new ConflictException(`Выбранная техника уже занята ${this.dateKey(workDate)}`);
2249: const assignedForDate = sourceVehicles
2251:   .filter((vehicle) => !occupied.has(this.planSlotKey(vehicle.id, workDate, 'day')))
2254:   .slice(0, demand.requiredCount);
2264: await tx.equipmentPlanAssignment.create({ data: { vehicleId: vehicle.id, workDate, shift: 'day' }});
        """,
    )
    add_code(
        doc,
        "6.5. Поиск замены при сервисном риске",
        r"D:\test-flow-app\backend\src\equipment-plans\equipment-plans.service.ts:1134",
        """
1134: private async findReplacementOptionsForServiceRisk(params) {
1142:   this.prisma.fleetVehicle.findMany({
1144:     id: { not: params.event.vehicleId },
1145:     type: params.event.vehicle.type,
1146:     status: { in: ['active', 'reserve'] },
1190:   return vehicles.filter((vehicle) =>
1192:     params.affectedPlans.every((plan) => {
1198:       const hasServiceBlock = blockingServices.some((event) =>
1201:         this.serviceOverlapsDate(event, plan.workDate));
1203:       return !occupiedSlots.has(slotKey) && !hasServiceBlock;
1207:   }).map((vehicle) => this.formatDraftVehicle(vehicle));
        """,
    )
    add_code(
        doc,
        "6.6. Разрешение риска: переназначение или сдвиг",
        r"D:\test-flow-app\backend\src\equipment-plans\equipment-plans.service.ts:1330",
        """
1330: if (dto.action === ServiceRiskResolutionAction.replace) {
1339:   const selectedReplacement = replacementOptions.find((vehicle) =>
1340:     vehicle.id === dto.replacementVehicleId);
1348:   const result = await this.prisma.equipmentPlanAssignment.updateMany({
1352:     vehicleId: selectedReplacement.id,
1353:     notes: `Переназначено из-за сервисного риска "${event.title}".`,
1380: const shiftSuggestion = this.buildServiceRiskShiftSuggestion(event, affectedPlans);
1482: await tx.roadWorkStage.update({
1485:   startDate: this.addDays(stage.startDate, shiftSuggestion.shiftDays),
1493: await tx.equipmentPlanAssignment.update({
1496:   workDate: this.addDays(plan.workDate, shiftSuggestion.shiftDays),
        """,
    )
    add_code(
        doc,
        "6.7. Уникальность слота в базе данных",
        r"D:\test-flow-app\backend\prisma\schema.prisma:438",
        """
438: model EquipmentPlanAssignment {
443:   vehicleId String
444:   workDate  DateTime
445:   shift     EquipmentPlanShift
459:   @@unique([vehicleId, workDate, shift])
        """,
    )
    add_code(
        doc,
        "6.8. Закрытые сервисные заявки не попадают в карточку рисков",
        r"D:\test-flow-app\next-app\app\dashboard\planning\page.tsx:1669",
        """
1669: const serviceRisks = serviceEvents.filter((event) => {
1670:   if (event.status === "completed") return false
1671:   if (!plannedVehicleIds.has(event.vehicleId)) return false
1672:   if (event.status === "in_progress" || event.status === "overdue") return true
1676:   const serviceRange = getServiceEventRange(event)
1682:   return serviceRange.end >= riskStart && serviceRange.start <= riskEnd
        """,
    )

    doc.add_heading("7. Результаты тестирования", level=1)
    test_rows = [
        ["Автосдвиг при занятости", "equipment-plans.service.spec.ts", "Этап переносится с 10.06.2026 на 11.06.2026, следующий этап - на 12.06.2026", "Passed"],
        ["ТО/ремонт в период работ", "equipment-plans.service.spec.ts", "Сервисное окно делает технику недоступной для автопланирования", "Passed"],
        ["Перекрытие внутри черновика", "equipment-plans.service.spec.ts", "Вторая перекрывающаяся стадия получает conflictCount=1", "Passed"],
        ["Перерасчет того же объекта", "equipment-plans.service.spec.ts", "replaceExisting исключает из конфликтов собственные старые назначения", "Passed"],
        ["Ручной выбор занятой техники", "equipment-plans.service.spec.ts", "Backend выбрасывает ConflictException", "Passed"],
        ["Гонка уникального слота БД", "equipment-plans.service.spec.ts", "P2002 переводится в понятную ошибку планирования", "Passed"],
        ["Сервисный риск с заменой", "equipment-plans.service.spec.ts", "Возвращается recommendedAction='replace' и варианты техники", "Passed"],
        ["Сервисный риск со сдвигом", "equipment-plans.service.spec.ts", "Сдвигаются даты этапа и сменные назначения", "Passed"],
        ["Статус техники при ремонте/ТО", "service-events.service.spec.ts", "Техника переходит в repair/maintenance и восстанавливает предыдущий статус", "Passed"],
    ]
    add_table(doc, ["Сценарий", "Тест", "Проверяемое поведение", "Статус"], test_rows, [2350, 2300, 3750, 960])
    doc.add_paragraph(
        "Фактический запуск: `npm test -- equipment-plans.service.spec.ts --runInBand` - 28 тестов прошли; "
        "`npm test -- service-events.service.spec.ts --runInBand` - 4 теста прошли."
    )
    add_code(
        doc,
        "7.1. Лог запуска тестов планирования",
        r"D:\test-flow-app\backend",
        """
> backend@0.0.1 test
> jest equipment-plans.service.spec.ts --runInBand

Test Suites: 1 passed, 1 total
Tests:       28 passed, 28 total
Snapshots:   0 total
Ran all test suites matching equipment-plans.service.spec.ts.

> backend@0.0.1 test
> jest service-events.service.spec.ts --runInBand

Test Suites: 1 passed, 1 total
Tests:       4 passed, 4 total
Snapshots:   0 total
Ran all test suites matching service-events.service.spec.ts.
        """,
    )

    doc.add_heading("8. Замечания критического агента и внесенные правки", level=1)
    critic_rows = [
        ["Не путать текущую версию с удаленными модулями", "В отчете оставлены только роли admin/user и актуальные разделы"],
        ["Не называть эвристику строгой оптимизацией", "Используется формулировка: эвристический подбор ближайшего доступного окна"],
        ["Показать отрицательные сценарии", "В матрицу включены занятость, ремонт, дефицит, невалидный этап, БД-конфликт"],
        ["Доказать влияние ремонта на план", "Добавлены скриншоты риска с заменой и риска со сдвигом, а также code excerpts"],
        ["Приложить тесты", "Добавлена таблица тестов и фактический результат запуска"],
    ]
    add_table(doc, ["Критическое замечание", "Как учтено в отчете"], critic_rows, [3800, 5560])

    doc.add_heading("9. Ограничения текущей реализации", level=1)
    limitations = [
        "Планирование является эвристическим: система подбирает ближайшие доступные даты и технику, но не решает задачу математической оптимизации стоимости или маршрутов.",
        "Нет GPS-телеметрии в реальном времени; мониторинг строится по статусам техники, планам и сервисным заявкам.",
        "Экономические показатели, стоимость машино-часа и расчет сметы намеренно не включены, так как тема диплома сфокусирована на планировании и мониторинге использования техники.",
        "Сценарии ремонта показывают эксплуатационный риск: если замена есть, система предлагает переназначение; если замены нет, срок работ увеличивается сдвигом этапов.",
    ]
    for item in limitations:
        doc.add_paragraph(item)

    doc.add_heading("10. Вывод", level=1)
    doc.add_paragraph(
        "Реализованная функция планирования соответствует теме диплома: она связывает дорожный объект, технологические этапы, типы строительной техники, "
        "сменные назначения, ТО/ремонты и риски срыва сроков. Главная инженерная составляющая находится на backend-уровне: расчет потребности, "
        "проверка занятости, учет ремонтных окон, запрет двойного назначения и автоматический сдвиг графика. Диаграмма Ганта и карточки рисков "
        "на frontend являются визуализацией этих расчетов."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)


if __name__ == "__main__":
    build()
